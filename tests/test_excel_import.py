"""Importing an old hand-kept budget workbook, built here from made-up numbers, then a bank export taking over."""
import datetime as dt

import openpyxl
import pytest
from openpyxl.styles import PatternFill

from budget import personal
from budget.csv_import import parse_csv, replace_spreadsheet_rows, store_transactions
from budget.excel_import import import_workbook

from conftest import account, category

SETTINGS = """
[spreadsheet]
ledger_sheet = "Budget"
name_sheets = ["Card Export"]
name_column = "G"
investments_sheet = "Investments"

[spreadsheet.ledger_accounts]
"Checking Income" = { account = "Checking", kind = "checking" }
"Checking Cost" = { account = "Checking", kind = "checking" }
"Credit Cost" = { account = "Credit Card", kind = "credit" }

[spreadsheet.balance_columns]
"Checking Balance" = { account = "Checking", kind = "checking" }
"Card Balance" = { account = "Credit Card", kind = "credit" }

[spreadsheet.panel_accounts]
"401k" = { account = "401k", kind = "retirement" }
"Home Loan" = { account = "Mortgage", kind = "loan" }

[spreadsheet.category_labels]
"Food" = { category = "Dining & Drinks", also_ok = ["Groceries"] }
"Misc" = { category = "Other", only_if_uncategorized = true }

[spreadsheet.investments]
"Brokerage" = { account = "Brokerage", kind = "brokerage", column = "E" }
"""

THROUGH = dt.date(2026, 2, 28)
YELLOW = PatternFill("solid", fgColor="FFFFFF00")
D = lambda day: dt.datetime.fromisoformat(day)


def build_workbook(path):
    """Columns: A date | B,C checking income | D,E checking cost | F,G card cost | H, I running balances."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Jan 2026 - Dec 2026", "Checking Income | Item", None, "Checking Cost | Item", None,
               "Credit Cost | Item", None, "Checking Balance", "Card Balance"])
    ws.append([])
    ws.append([])
    rows = [
        (4, "2026-01-02", {"B": -3000, "C": "ACME PAYROLL"}),
        (5, "2026-01-05", {"D": 54.21, "E": "Safeway"}),
        (6, "2026-01-07", {"F": 23.50, "G": "Food Cart"}),
        (7, "2026-01-09", {"F": 40, "G": "Mystery", "D": 12.00, "E": "Safeway"}),
        (8, "2026-01-12", {"F": 18, "G": "Corner Cafe"}),
        (9, "2026-01-14", {"F": 22, "G": "Corner Cafe"}),
        (10, "2026-01-16", {"F": 31, "G": "Corner Cafe"}),
        (11, "2026-01-18", {"F": 9.75, "G": "Bagel Place"}),   # yellow like the Food cells, no formula
        (12, "2026-01-31", {"H": 5000, "I": -63.50}),
        (13, "2026-02-03", {"D": 100, "E": "Plumber"}),
        (14, "2026-02-27", {"H": 4900, "I": -10}),
        (15, "2026-03-05", {"D": 999, "E": "Planned bill"}),  # after the workbook's through date
    ]
    for r, day, cells in rows:
        ws[f"A{r}"] = D(day)
        for col, value in cells.items():
            ws[f"{col}{r}"] = value
    for r in (6, 8, 9, 10, 11):
        ws[f"F{r}"].fill = YELLOW
    # Summaries right of the ledger: a label with a =SUM(...) of the ledger cells it covers.
    ws["Y4"], ws["Y5"] = "Food", "=SUM(F6,F8:F10,D5)"
    ws["Z4"], ws["Z5"] = "Misc", "=SUM(F7,D7)"
    # Label/value panel: the value sits left of its label, in the month block it's in.
    ws["AB5"], ws["AC5"] = 45_000, "401k"
    ws["AB13"], ws["AC13"] = 46_250.75, "401k"
    ws["AB14"] = '=-FV(0.06/12,DATEDIF(43845,A14,"m"),PMT(0.06/12,30*12,200000),200000)'
    ws["AC14"] = "Home Loan"

    export = wb.create_sheet("Card Export")
    export.append(["Date", "Transaction", "Name", "Memo", "Amount", None, "My name"])
    for i in range(3):
        export.append([f"2026-01-{10 + i}", "DEBIT", "SQ *TACO TRUCK 1234   SPRINGFIELD CO", "2469; 05812; ; ;", -9, None, "Taco Truck"])

    inv = wb.create_sheet("Investments")
    inv.append(["Brokerage"])
    inv.append(["Fund", None, None, None, "Value"])
    inv.append(["Index fund", None, None, None, 12_000.50])
    inv.append(["Bond fund", None, None, None, "$3,000.25"])
    wb.save(path)
    return path


@pytest.fixture
def setup(app, conn, tmp_path):
    config = tmp_path / "data" / "personal.toml"
    config.write_text(SETTINGS)
    path = build_workbook(tmp_path / "data" / "source" / "budget.xlsx")
    return conn, path, personal.load(config)


def rows(conn, account_name):
    return {(r["date"], r["amount"]): dict(r) for r in conn.execute(
        """SELECT t.date, t.amount, t.name, t.category_source, t.notes, c.name AS category
           FROM transactions t JOIN accounts a ON a.id = t.account_id LEFT JOIN categories c ON c.id = t.category_id
           WHERE a.name = ?""", (account_name,))}


def balances(conn, account_name):
    return {r[0]: (r[1], r[2]) for r in conn.execute(
        "SELECT b.month, b.amount, b.as_of FROM balances b JOIN accounts a ON a.id = b.account_id WHERE a.name = ?",
        (account_name,))}


def test_ledger_rows_and_signs(setup):
    conn, path, settings = setup
    result = import_workbook(conn, path, THROUGH, settings)
    assert result["through"] == "2026-02-28"
    assert dict((a, (read, added)) for a, read, added in result["ledger"]) == {"Checking": (4, 4), "Credit Card": (6, 6)}
    checking = rows(conn, "Checking")
    assert set(checking) == {("2026-01-02", 3000.0), ("2026-01-05", -54.21), ("2026-01-09", -12.0), ("2026-02-03", -100.0)}
    assert checking[("2026-01-02", 3000.0)]["category"] == "Paycheck"
    assert ("2026-03-05", -999.0) not in checking  # planned, after the save date
    assert account_kind(conn, "Credit Card") == "credit"


def account_kind(conn, name):
    return conn.execute("SELECT kind FROM accounts WHERE name = ?", (name,)).fetchone()[0]


def test_hand_sorted_categories(setup):
    conn, path, settings = setup
    result = import_workbook(conn, path, THROUGH, settings)
    card, checking = rows(conn, "Credit Card"), rows(conn, "Checking")
    # Food: no rule knew the cart, so the spreadsheet's bucket wins; Safeway's Groceries is also fine for Food.
    assert (card[("2026-01-07", -23.5)]["category"], card[("2026-01-07", -23.5)]["category_source"]) == ("Dining & Drinks", "sheet")
    assert checking[("2026-01-05", -54.21)]["category"] == "Groceries"
    assert checking[("2026-01-05", -54.21)]["category_source"] == "rule"
    # Misc only fills in what nothing else categorized.
    assert (card[("2026-01-09", -40.0)]["category"], card[("2026-01-09", -40.0)]["category_source"]) == ("Other", "sheet")
    assert checking[("2026-01-09", -12.0)]["category"] == "Groceries"
    # A cell with no formula pointing at it takes the bucket its color means nearby.
    assert card[("2026-01-18", -9.75)]["category"] == "Dining & Drinks"
    assert result["hand_sorted"] == {"formula": 7, "color": 1}


def test_balances_loans_names_and_investments(setup):
    conn, path, settings = setup
    result = import_workbook(conn, path, THROUGH, settings)
    assert balances(conn, "Checking") == {"2026-01": (5000.0, "2026-01-31"), "2026-02": (4900.0, "2026-02-27")}
    assert balances(conn, "Credit Card") == {"2026-01": (63.5, "2026-01-31"), "2026-02": (10.0, "2026-02-27")}
    assert balances(conn, "401k") == {"2026-01": (45_000.0, "2026-01-31"), "2026-02": (46_250.75, "2026-02-28")}
    assert balances(conn, "Brokerage") == {"2026-02": (15_000.75, "2026-02-28")}
    loan = dict(conn.execute("SELECT * FROM loan_terms").fetchone())
    assert (loan["principal"], loan["annual_rate"], loan["term_months"]) == (200_000, 6, 360)
    assert (loan["start_date"], loan["counts_from"]) == ("2020-01-15", "2026-02-01")
    assert result["loans_added"] == 1 and account_kind(conn, "Mortgage") == "loan"
    rule = conn.execute("SELECT rename_to, source FROM rules WHERE pattern = 'TACO TRUCK'").fetchone()
    assert tuple(rule) == ("Taco Truck", "excel")


def test_reimport_changes_nothing_you_edited(setup):
    conn, path, settings = setup
    import_workbook(conn, path, THROUGH, settings)
    count = conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0]
    cart = conn.execute("SELECT id FROM transactions WHERE name = 'Food Cart'").fetchone()[0]
    conn.execute("UPDATE transactions SET category_id = ?, category_source = 'manual' WHERE id = ?",
                 (category(conn, "Travel"), cart))
    conn.execute("UPDATE balances SET amount = 4321, source = 'manual' WHERE month = '2026-02' "
                 "AND account_id = (SELECT id FROM accounts WHERE name = 'Checking')")
    conn.execute("UPDATE loan_terms SET annual_rate = 5.5")
    result = import_workbook(conn, path, THROUGH, settings)
    assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == count
    assert [added for _, _, added in result["ledger"]] == [0, 0] and result["rules_added"] == 0
    assert conn.execute("SELECT category_id FROM transactions WHERE id = ?", (cart,)).fetchone()[0] == category(conn, "Travel")
    assert balances(conn, "Checking")["2026-02"] == (4321.0, "2026-02-27")
    assert conn.execute("SELECT annual_rate FROM loan_terms").fetchone()[0] == 5.5 and result["loans_added"] == 0


BANK = """Date,Transaction,Name,Memo,Amount
2026-01-04,CREDIT,ACME CORP PAYROLL,,3000.00
2026-01-06,DEBIT,SAFEWAY #1234 SPRINGFIELD CO,24692163092; 05411; ; ; ;,-54.21
2026-01-20,DEBIT,HARDWARE STORE,,-15.00
"""


def test_bank_export_replaces_spreadsheet_rows(setup):
    conn, path, settings = setup
    import_workbook(conn, path, THROUGH, settings)
    checking = account(conn, "Checking")
    conn.execute("UPDATE transactions SET notes = 'for the party', date_override = '2026-01-05' "
                 "WHERE name = 'Safeway' AND amount = -54.21")
    parsed = parse_csv(BANK)
    store_transactions(conn, checking, "bank.csv", parsed)
    removed, carried = replace_spreadsheet_rows(conn, checking, parsed, "bank.csv")
    # Jan 1-20 is the bank's now: the sheet's paycheck, Safeway and the 9th's Safeway go; February stays.
    assert removed == 3 and carried == 1
    got = rows(conn, "Checking")
    assert set(got) == {("2026-01-04", 3000.0), ("2026-01-06", -54.21), ("2026-01-20", -15.0), ("2026-02-03", -100.0)}
    safeway = conn.execute("SELECT notes, date_override, effective_date FROM transactions WHERE date = '2026-01-06'").fetchone()
    assert tuple(safeway) == ("for the party", "2026-01-05", "2026-01-05")
    # The sheet's running balances from the bank's start on go too (typed, never checked against the bank).
    assert balances(conn, "Checking") == {} and len(balances(conn, "Credit Card")) == 2

    # Importing the workbook again doesn't bring the replaced rows back.
    result = import_workbook(conn, path, THROUGH, settings)
    assert set(rows(conn, "Checking")) == set(got)
    assert result["skipped_for_bank_data"] == 3
    assert balances(conn, "Checking") == {}
    # Nor does uploading the same export again pair or remove anything.
    store_transactions(conn, checking, "bank.csv", parsed)
    assert replace_spreadsheet_rows(conn, checking, parsed, "bank.csv") == (0, 0)
    assert set(rows(conn, "Checking")) == set(got)


def test_not_configured(conn, tmp_path):
    path = build_workbook(tmp_path / "plain.xlsx")
    result = import_workbook(conn, path, THROUGH, personal.Personal())
    assert result["configured"] is False
    assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 0


def test_upload_route_reads_personal_toml(setup, client):
    conn, path, _ = setup
    r = client.post("/upload/excel", data={"existing": path.name, "through": "2026-02-28"}, follow_redirects=True)
    assert r.status_code == 200 and "Imported budget.xlsx through 2026-02-28" in r.get_data(as_text=True)
    assert conn.execute("SELECT COUNT(*) FROM transactions").fetchone()[0] == 10


def test_upload_page_offers_the_workbook_import_once_configured(setup, client):
    html = client.get("/upload").get_data(as_text=True)
    assert "Import workbook" in html and "Workbook in data/source" in html
    assert "Have an old budget spreadsheet?" not in html
