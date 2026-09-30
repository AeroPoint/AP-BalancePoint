"""Blackjack: reading a tracker workbook, sessions in the bankroll account, and the log form."""
import datetime as dt

import openpyxl
import pytest

from budget import blackjack as bj
from budget.balances import load_ledgers


@pytest.fixture
def workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "2025 Winnings Tracker"
    ws.append(["Location", "Date", "EV Total", "Tip", "Hrs", "EV Per Hour", "AV Total", "AV Per Hour", 17, "DAS", "BJ",
               "Decks", "Spread", "Min", "Max", "Miles Of Travel", "Room n Board", "Notes"])
    ws.append(["Riverside", "Mar 1 2025", 50, 0, dt.time(1, 0), 50, 300, None, "H", "Y", "3:2", 6, "1-12", 15, 300, 40, 0, "fun"])
    ws.append([None, None, 25, 0, dt.time(0, 30), 50, None, None, "S", "N", "3:2", 2, "1-8", 25, 500])  # 2nd table
    ws.append(["Canyon Club", "Mar 2 2025", 40, 5, dt.time(2, 0), 20, -150, None, "H", "Y", dt.time(3, 2), 6,
               dt.datetime(2025, 1, 20), 25, 500, 0, 80])  # Excel read "3:2" as 3:02 and "1-20" as a date
    ws.append(["ALL TIME", None, 115, None, None, None, 150])
    r = wb.create_sheet("Weekend trip")
    r.append([None, "Location", "Decks", "Pen", "EV", None, "Other"])
    r.append(["Called", "Canyon Club", 6, 1, 41.0, 20.5, "friendly"])
    r.append([None, None, 2, 0.6, 88.0, None, None])  # another game at the same casino
    r.append([None, "Average of Trip", 4, 0.8, 64.5])
    r.append([None, None, None, None, "Spread A"])
    r.append([None, None, None, None, "TC", "Bet", "NumHand"])
    r.append([None, None, None, None, 1, 25, 1])
    r.append([None, None, None, None, 2, 50, 2])
    t = wb.create_sheet("2025 Training Tracker")
    t.append([None, None, None, "Testout", "Count"])
    t.append([dt.datetime(2025, 2, 6), dt.time(0, 45), "Drills", "2D", "Hi-Lo"])
    path = tmp_path / "tracker.xlsx"
    wb.save(path)
    return path


def test_read_workbook(workbook):
    d = bj.read_workbook(workbook)
    assert [(s["location"], s["result"], len(s["tables"])) for s in d["sessions"]] == [("Riverside", 300, 2), ("Canyon Club", -150, 1)]
    canyon = d["sessions"][1]["tables"][0]["rules"]
    assert canyon["BJ"] == "3:2" and canyon["Spread"] == "1-20"
    assert d["sessions"][1]["tip"] == 5 and d["sessions"][0]["miles"] == 40
    assert [r["casino"] for r in d["research"]] == ["Canyon Club", "Canyon Club"]
    assert d["research"][0]["fields"]["EV (B)"] == 20.5
    notes = {n["title"]: n["body"] for n in d["notes"]}
    assert "Average of Trip" in notes and notes["Spread A"].splitlines()[1] == "1 | 25 | 1"
    assert d["training"][0]["minutes"] == 45


def test_import_keeps_sessions_logged_in_the_app(conn, workbook):
    acct = bj.bankroll_account(conn)
    got = bj.import_workbook(conn, workbook, acct)
    assert (got["sessions"], got["tables"], got["net"]) == (2, 3, 150)
    bj.save_session(conn, {"date": "2025-03-05", "location": "Lakeview", "result": 80,
                           "tables": [{"hours": 1.5, "ev_hour": 40, "rules": {"Decks": 2}}]}, acct)
    bj.import_workbook(conn, workbook, acct)  # again
    sessions = bj.sessions(conn)
    assert [s["location"] for s in sessions] == ["Riverside", "Canyon Club", "Lakeview"]
    total = conn.execute("SELECT SUM(amount) FROM transactions WHERE account_id = ?", (acct,)).fetchone()[0]
    assert total == 230  # every session's result is in the bankroll account, once
    lake = sessions[-1]
    assert (lake["hours"], lake["ev"], lake["ev_hour"]) == (1.5, 60.0, 40.0)


def test_log_form_round_trip(conn, client):
    form = {"date": "2026-01-10", "location": "Test Casino", "outcome": "lost", "result": "1,200",
            "table0-hours": "1:30", "table0-ev_total": "90", "table0-Decks": "6", "table0-17": "S",
            "table1-hours": "0:30", "table1-ev_hour": "60", "table1-Decks": "2"}
    assert client.post("/bankroll/sessions/save", data=form).status_code == 302
    s = bj.sessions(conn)[-1]
    assert (s["result"], len(s["tables"]), s["hours"], s["ev"]) == (-1200, 2, 2.0, 120.0)
    assert bj.rules_line(s["tables"][0]["rules"]) == "6D · S17"
    acct = bj.bankroll_account(conn)
    assert load_ledgers(conn)[acct].has_transactions
    client.post(f"/bankroll/sessions/{s['id']}/delete")
    assert conn.execute("SELECT COUNT(*) FROM transactions WHERE account_id = ?", (acct,)).fetchone()[0] == 0
