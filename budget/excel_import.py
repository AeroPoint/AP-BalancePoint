"""Import of an old budget workbook, driven by the [spreadsheet] section of data/personal.toml.

Running it again skips what's already in. Depending on what the settings describe, it pulls:
  1. Merchant names - sheets holding a US Bank export next to the names you typed for each charge.
  2. Transaction history - a daily ledger sheet whose top rows are a legend of which columns were
     used from which month ("Jan 2026 - Dec 2026 | Checking Income | Item ...").
  3. Hand-sorted categories - per-month "=SUM(F68,H70,...)" formulas under labels mapped to
     categories, with cell colors filling in for ledger cells no formula points at.
  4. Month-end balances - running balance columns named in the legend, plus label/value pairs
     elsewhere on the ledger sheet.
  5. Current values from an investments sheet, one titled section per account.
"""
import bisect
import datetime as dt
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

from . import seed
from .balances import month_end
from .csv_import import bank_start, mcc_from_memo, store_transactions
from .personal import PersonalConfigError
from .rules import RuleEngine, merchant_key

MONTHS = {m: i for i, m in enumerate(["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}
IGNORED_FRIENDLY = {"PAID", "#N/A", "ERROR", "-", ""}
CELL_REF = re.compile(r"\$?([A-Z]{1,3})\$?(\d+)(?::\$?([A-Z]{1,3})\$?(\d+))?")
# One loan inside a balance formula: FV(rate/12, DATEDIF(start, date, "m"), PMT(rate/12, term, amount), amount)
FV_LOAN = re.compile(
    r"FV\(\s*([\d.]+)\s*/\s*12\s*,\s*DATEDIF\(\s*(\d+)\s*,[^)]*\)\s*,\s*"
    r"PMT\(\s*[\d.]+\s*/\s*12\s*,\s*([\d.*\s]+?)\s*,\s*([\d.]+)\s*\)",
    re.IGNORECASE,
)


def _label(text):
    return " ".join(str(text).lower().split())


def settings(personal):
    """Normalize the [spreadsheet] table from personal.toml."""
    raw = personal.spreadsheet if personal else {}

    def account(entry, default_kind):
        return entry["account"], entry.get("kind", default_kind)

    try:
        labels = {}
        for key, entry in raw.get("category_labels", {}).items():
            wanted = seed.category_name(entry["category"])
            keep = None if entry.get("only_if_uncategorized") else {wanted, *map(seed.category_name, entry.get("also_ok", []))}
            labels[_label(key)] = (wanted, keep)
        cfg = {
            "ledger_sheet": raw.get("ledger_sheet", "Budget"),
            "name_sheets": list(raw.get("name_sheets", [])),
            "name_column": raw.get("name_column", "G"),
            "investments_sheet": raw.get("investments_sheet", "Investments"),
            "ledger_accounts": {k.strip(): account(v, "checking") for k, v in raw.get("ledger_accounts", {}).items()},
            "balance_columns": {_label(k): account(v, "checking") for k, v in raw.get("balance_columns", {}).items()},
            "panel_accounts": {_label(k): account(v, "other") for k, v in raw.get("panel_accounts", {}).items()},
            "category_labels": labels,
            "investments": {
                k.strip(): (*account(v, "brokerage"), v["column"]) for k, v in raw.get("investments", {}).items()
            },
        }
    except (KeyError, TypeError) as exc:
        raise PersonalConfigError(f"personal.toml [spreadsheet]: an entry is missing or malformed ({exc}).") from exc
    cfg["configured"] = any(cfg[k] for k in ("name_sheets", "ledger_accounts", "panel_accounts", "investments"))
    return cfg


def _fill_key(cell):
    fill = cell.fill
    if fill is None or fill.fill_type is None:
        return None
    color = fill.fgColor
    if color.type == "rgb":
        return None if color.rgb in ("00000000", "FFFFFFFF") else color.rgb
    if color.type == "theme":
        return f"theme{color.theme}{color.tint:+.2f}"
    if color.type == "indexed":
        return f"indexed{color.indexed}"
    return None


def _formula_cells(formula):
    """(row, 0-based column) for every cell a formula references, ranges expanded."""
    for m in CELL_REF.finditer(formula):
        c1, r1 = column_index_from_string(m.group(1)), int(m.group(2))
        c2 = column_index_from_string(m.group(3)) if m.group(3) else c1
        r2 = int(m.group(4)) if m.group(4) else r1
        for r in range(min(r1, r2), max(r1, r2) + 1):
            for c in range(min(c1, c2), max(c1, c2) + 1):
                yield r, c - 1


def ensure_account(conn, name, kind):
    row = conn.execute("SELECT id FROM accounts WHERE name = ?", (name,)).fetchone()
    if row:
        return row["id"]
    return conn.execute("INSERT INTO accounts (name, kind) VALUES (?, ?)", (name, kind)).lastrowid


def set_balance(conn, account_id, month, amount, as_of, source="excel"):
    """Upsert one month's balance. Balances typed into the app are never overwritten by an import."""
    return conn.execute(
        """INSERT INTO balances (account_id, month, amount, as_of, source) VALUES (?, ?, ?, ?, ?)
           ON CONFLICT (account_id, month) DO UPDATE
             SET amount = excluded.amount, as_of = excluded.as_of, source = excluded.source
           WHERE balances.source != 'manual'""",
        (account_id, month, amount, as_of, source),
    ).rowcount


# ---------------------------------------------------------------- merchant names

def _plausible(key, friendly):
    """The name sheets may pair names by dollar amount, so a rarely-seen pair must share a word
    (or initials, like BWW) with the bank text before we trust it."""
    words = re.findall(r"[a-z0-9]+", key.lower())
    compact, initials = "".join(words), "".join(w[0] for w in words)
    return any(
        len(tok) >= 3 and (tok[:4] in compact or tok in initials)
        for tok in re.findall(r"[a-z0-9]+", friendly.lower())
    )


def import_dictionary(conn, wb, cfg):
    """Build 'raw merchant -> your name' rules by majority vote over the name sheets."""
    friendly_col = column_index_from_string(cfg["name_column"]) - 1
    names, mccs = defaultdict(Counter), defaultdict(Counter)
    for sheet in cfg["name_sheets"]:
        if sheet not in wb.sheetnames:
            continue
        for row in wb[sheet].iter_rows(min_row=2, max_col=max(friendly_col + 1, 4), values_only=True):
            raw, memo, friendly = row[2], row[3], row[friendly_col]
            if not raw or not isinstance(friendly, str) or friendly.strip().upper() in IGNORED_FRIENDLY:
                continue
            key = merchant_key(raw)
            if len(key) < 3:
                continue
            names[key][friendly.strip()] += 1
            if mcc := mcc_from_memo(memo):
                mccs[key][mcc] += 1

    cat_ids = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM categories")}
    engine = RuleEngine(conn)
    added = 0
    for key, counts in names.items():
        friendly, top = counts.most_common(1)[0]
        total = sum(counts.values())
        if top / total < 0.6 or (total >= 3 and top < 2):
            continue
        if top < 3 and not _plausible(key, friendly):
            continue
        # A dictionary/keyword category (BWW -> Eating Out) beats the card-type guess (5813 = bar).
        known = engine.resolve(key, fixed_name=friendly)
        if known.source == "rule":
            category_id = known.category_id
        else:
            category_id = cat_ids.get(seed.mcc_category(mccs[key].most_common(1)[0][0])) if mccs[key] else None
        existing = conn.execute(
            "SELECT id, source FROM rules WHERE match_on = 'raw' AND pattern = ?", (key,)
        ).fetchone()
        if existing is None:
            conn.execute(
                "INSERT INTO rules (match_on, pattern, rename_to, category_id, source) VALUES ('raw', ?, ?, ?, 'excel')",
                (key, friendly, category_id),
            )
            added += 1
        elif existing["source"] in ("builtin", "config"):
            # The names you typed in the spreadsheet win over default names for the same merchant.
            conn.execute(
                "UPDATE rules SET rename_to = ?, source = 'excel', category_id = COALESCE(category_id, ?) WHERE id = ?",
                (friendly, category_id, existing["id"]),
            )
            added += 1
    return added


# ---------------------------------------------------------------- ledger sheet

def _legend(ws, cfg):
    """[(first month, {ledger col: header}, {balance col: account})], newest layout first."""
    legend = []
    for row in ws.iter_rows(min_row=1, max_row=15, max_col=45):
        since = re.match(r"\s*([A-Za-z]{3})[a-z]* (\d{4})", str(row[0].value or ""))
        if not since or since.group(1).title() not in MONTHS:
            continue
        ledger, balances = {}, {}
        for i, cell in enumerate(row):
            if not isinstance(cell.value, str):
                continue
            if "| Item" in cell.value:
                ledger[i] = cell.value.split("|")[0].strip()
            elif _label(cell.value) in cfg["balance_columns"]:
                balances[i] = cfg["balance_columns"][_label(cell.value)]
        legend.append((dt.date(int(since.group(2)), MONTHS[since.group(1).title()], 1), ledger, balances))
    legend.sort(key=lambda item: item[0], reverse=True)
    return legend


def _layout_for(legend, day):
    return next(((ledger, balances) for start, ledger, balances in legend if start <= day), (None, None))


def _hand_sorted(values_ws, formulas_ws, legend, labels):
    """{(row, col): (label, 'formula' | 'color')} for ledger amount cells you sorted into a bucket."""
    refs, sizes = defaultdict(set), Counter()
    for row in formulas_ws.iter_rows(min_row=4, min_col=23):  # summaries live right of the ledger
        for cell in row:
            if not isinstance(cell.value, str) or cell.value.startswith("="):
                continue
            key = _label(cell.value)
            if key not in labels:
                continue
            below = formulas_ws.cell(cell.row + 1, cell.column).value
            if isinstance(below, str) and below.startswith("="):
                for ref in _formula_cells(below):
                    refs[ref].add(key)
                    sizes[key] += 1

    def pick(found):
        # A label that only fills in uncategorized rows loses to any more specific label.
        specific = [label for label in found if labels[label][1] is not None] or list(found)
        return min(specific, key=lambda label: (sizes[label], label))

    sorted_cells, colors, unsorted = {}, defaultdict(Counter), []
    for row in values_ws.iter_rows(min_row=4, max_col=24):
        first = row[0].value
        if not isinstance(first, dt.datetime):
            continue
        ledger, _ = _layout_for(legend, first.date())
        month = first.year * 12 + first.month
        for i in ledger or {}:
            if i + 1 >= len(row) or not isinstance(row[i].value, (int, float)):
                continue
            position = (row[0].row, i)
            fill = _fill_key(row[i + 1]) or _fill_key(row[i])
            if position in refs:
                label = pick(refs[position])
                sorted_cells[position] = (label, "formula")
                if fill:
                    colors[(fill, month)][label] += 1
            elif fill:
                unsorted.append((position, fill, month))

    # Colors get reused for different buckets over the years, so a color's meaning comes from
    # formula-sorted cells within three months either side, and only when nearly all of them agree.
    for position, fill, month in unsorted:
        counts = Counter()
        for offset in range(-3, 4):
            counts.update(colors.get((fill, month + offset), {}))
        if not counts:
            continue
        label, top = counts.most_common(1)[0]
        if labels[label][1] is not None and top >= 3 and top / sum(counts.values()) >= 0.8:
            sorted_cells[position] = (label, "color")
    return sorted_cells


def _bank_windows(conn):
    """{account name: [(first day, last day)]} of the dates bank exports are the record for."""
    windows = defaultdict(list)
    for name, since, first, last in conn.execute(
        """SELECT a.name, a.bank_from, MIN(t.date), MAX(t.date) FROM transactions t
           JOIN imports i ON i.id = t.import_id JOIN accounts a ON a.id = t.account_id
           WHERE i.kind = 'csv' GROUP BY t.import_id"""
    ):
        windows[name].append((bank_start(since and f"{since}-01", first), last))
    return windows


def import_ledger(conn, values_ws, formulas_ws, filename, through, cfg):
    """Returns (summary dict, month-end balances {(account, kind): {month: amount}})."""
    legend = _legend(values_ws, cfg)
    labels = cfg["category_labels"]
    hand_sorted = _hand_sorted(values_ws, formulas_ws, legend, labels) if legend and labels else {}
    cat_ids = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM categories")}
    by_account, closing, matched = defaultdict(list), defaultdict(dict), Counter()
    # Bank exports are the record for the dates they cover (see csv_import.replace_spreadsheet_rows).
    covered = _bank_windows(conn)
    skipped = 0
    header = None
    for row in values_ws.iter_rows(max_col=45):
        first = row[0].value
        headers = {i: c.value.split("|")[0].strip() for i, c in enumerate(row) if isinstance(c.value, str) and "| Item" in c.value}
        if headers:
            header = headers
            continue
        if not isinstance(first, dt.datetime) or first.date() > through:
            continue
        day, r = first.date(), row[0].row
        ledger, balance_cols = _layout_for(legend, day)
        for i, column in (ledger or header or {}).items():
            amount = row[i].value if i < len(row) else None
            if column not in cfg["ledger_accounts"] or not isinstance(amount, (int, float)) or round(amount, 2) == 0:
                continue
            if any(start <= day.isoformat() <= end for start, end in covered[cfg["ledger_accounts"][column][0]]):
                skipped += 1
                continue
            item = row[i + 1].value if i + 1 < len(row) else None
            name = str(item).strip() if item not in (None, "") else "Unlabeled"
            # Sheet stores costs positive and income negative; the app uses negative = money out.
            entry = {
                "date": day.isoformat(),
                "amount": round(-amount, 2),
                "raw": name,
                "name": name,
                "memo": f"{filename} {values_ws.title}!{get_column_letter(i + 1)}{r}",
            }
            if (r, i) in hand_sorted:
                label_key, how = hand_sorted[(r, i)]
                wanted, compatible = labels[label_key]
                if wanted in cat_ids:
                    keep = {cat_ids[c] for c in compatible if c in cat_ids} if compatible else None
                    entry["category_hint"] = (cat_ids[wanted], keep)
                    matched[how] += 1
            by_account[cfg["ledger_accounts"][column]].append(entry)
        for i, target in (balance_cols or {}).items():
            value = row[i].value if i < len(row) else None
            if isinstance(value, (int, float)):  # rows run in date order, so the month's last row wins
                amount = -value if seed.account_side(target[1]) == "liability" else value
                closing[target][day.strftime("%Y-%m")] = (amount, day.isoformat())

    engine = RuleEngine(conn)
    summary, recategorized = [], 0
    for (account, kind), rows in by_account.items():
        account_id = ensure_account(conn, account, kind)
        _, read, added, hinted, _ = store_transactions(conn, account_id, f"{filename} ({account})", rows, "excel", engine)
        summary.append((account, read, added))
        recategorized += hinted
    return {"ledger": summary, "hand_sorted": dict(matched), "recategorized": recategorized,
            "skipped_for_bank_data": skipped}, closing


def _panel_balances(ws, through, panel):
    if not panel:
        return {}
    date_rows = [(c.row, c.value.date()) for (c,) in ws.iter_rows(min_row=4, max_col=1) if isinstance(c.value, dt.datetime)]
    row_numbers = [r for r, _ in date_rows]
    found = defaultdict(dict)
    for row in ws.iter_rows(min_row=4, min_col=24):
        for cell in row:
            target = panel.get(_label(cell.value)) if isinstance(cell.value, str) else None
            if not target:
                continue
            value = ws.cell(cell.row, cell.column - 1).value
            idx = bisect.bisect_left(row_numbers, cell.row)  # a panel belongs to the month block it sits in
            if not isinstance(value, (int, float)) or idx >= len(date_rows) or date_rows[idx][1].replace(day=1) > through:
                continue
            month = date_rows[idx][1].strftime("%Y-%m")
            amount = abs(value) if seed.account_side(target[1]) == "liability" else value
            found[target][month] = (amount, min(month_end(month), through.isoformat()))
    return found


def _loan_terms(formulas_ws, values_ws, through, panel):
    """Loan terms from panel cells that calculate a balance with =-FV(...) instead of a typed number.

    Returns {(account, kind): (terms, first day of the first month the formula was used)}; the
    calculation takes over from that month, so earlier typed-in balances stay as recorded.
    """
    if not panel:
        return {}
    date_rows = [(c.row, c.value.date()) for (c,) in values_ws.iter_rows(min_row=4, max_col=1) if isinstance(c.value, dt.datetime)]
    row_numbers = [r for r, _ in date_rows]
    found = {}
    for row in formulas_ws.iter_rows(min_row=4, min_col=24):
        for cell in row:
            target = panel.get(_label(cell.value)) if isinstance(cell.value, str) else None
            if not target:
                continue
            formula = formulas_ws.cell(cell.row, cell.column - 1).value
            idx = bisect.bisect_left(row_numbers, cell.row)
            if not isinstance(formula, str) or idx >= len(date_rows) or date_rows[idx][1].replace(day=1) > through:
                continue
            terms = []
            for m in FV_LOAN.finditer(formula):
                months = 1.0
                for factor in m.group(3).split("*"):
                    months *= float(factor)
                terms.append({
                    "annual_rate": round(float(m.group(1)) * 100, 6),
                    "start_date": (dt.date(1899, 12, 30) + dt.timedelta(days=int(m.group(2)))).isoformat(),
                    "term_months": int(round(months)),
                    "principal": float(m.group(4)),
                })
            if terms:
                first = found[target][1] if target in found else date_rows[idx][1].replace(day=1).isoformat()
                found[target] = (terms, first)  # rows run top to bottom: newest terms, first month used
    return found


def _write_loans(conn, loans):
    added = 0
    for (name, kind), (terms, counts_from) in loans.items():
        account_id = ensure_account(conn, name, kind)
        if conn.execute("SELECT 1 FROM loan_terms WHERE account_id = ?", (account_id,)).fetchone():
            continue  # already set up, possibly edited in the app: leave it alone
        for t in terms:
            conn.execute(
                """INSERT INTO loan_terms (account_id, label, principal, annual_rate, term_months, start_date, counts_from)
                   VALUES (?, ?, ?, ?, ?, ?, ?)""",
                (account_id, f"${t['principal']:,.0f} loan", t["principal"], t["annual_rate"],
                 t["term_months"], t["start_date"], counts_from),
            )
            added += 1
    return added


# ---------------------------------------------------------------- investments sheet

def _money(value):
    if isinstance(value, (int, float)):
        return float(value)
    digits = re.sub(r"[^\d.]", "", str(value or ""))
    try:
        return float(digits) if digits else 0.0
    except ValueError:
        return 0.0


def _investment_balances(wb, cfg):
    """Per-account totals from titled sections on the investments sheet."""
    sections = cfg["investments"]
    if not sections or cfg["investments_sheet"] not in wb.sheetnames:
        return {}
    current, totals = None, Counter()
    for row in wb[cfg["investments_sheet"]].iter_rows(values_only=True):
        filled = [v for v in row if v not in (None, "")]
        if len(filled) == 1 and isinstance(filled[0], str):
            current = sections.get(filled[0].strip())  # a title alone on its row starts a section
            continue
        if current:
            account, kind, column = current
            idx = column_index_from_string(column) - 1
            if idx < len(row):
                totals[(account, kind)] += _money(row[idx])
    return {target: round(amount, 2) for target, amount in totals.items() if amount}


def _write_balances(conn, found):
    written = 0
    for (name, kind), months in found.items():
        account_id = ensure_account(conn, name, kind)
        items = sorted(months.items())
        for idx, (month, (amount, as_of)) in enumerate(items):
            # A 0 between two real balances is a cell nobody updated that month, not an empty account.
            if round(amount, 2) == 0 and 0 < idx < len(items) - 1 and items[idx - 1][1][0] and items[idx + 1][1][0]:
                continue
            written += set_balance(conn, account_id, month, round(amount, 2), as_of)
    return written


def import_workbook(conn, path, through=None, personal=None):
    path = Path(path)
    cfg = settings(personal)
    result = {"configured": cfg["configured"], "through": None, "rules_added": 0,
              "ledger": [], "hand_sorted": {}, "recategorized": 0, "skipped_for_bank_data": 0, "balances": 0,
              "loans_added": 0}
    if not cfg["configured"]:
        return result
    values = openpyxl.load_workbook(path, data_only=True)
    # Rows after the workbook was last saved are projected bills, not real transactions. The save
    # date stored inside the workbook survives copying and uploading; the file's own date doesn't.
    saved = values.properties.modified
    through = through or (saved.date() if saved else dt.date.fromtimestamp(path.stat().st_mtime))
    result["through"] = through.isoformat()
    result["rules_added"] = import_dictionary(conn, values, cfg)
    found, loans = defaultdict(dict), {}
    if cfg["ledger_sheet"] in values.sheetnames and (cfg["ledger_accounts"] or cfg["panel_accounts"]):
        formulas = openpyxl.load_workbook(path)
        ledger_ws = values[cfg["ledger_sheet"]]
        summary, closing = import_ledger(conn, ledger_ws, formulas[cfg["ledger_sheet"]], path.name, through, cfg)
        result.update(summary)
        for source in (closing, _panel_balances(ledger_ws, through, cfg["panel_accounts"])):
            for target, months in source.items():
                found[target].update(months)
        loans = _loan_terms(formulas[cfg["ledger_sheet"]], ledger_ws, through, cfg["panel_accounts"])
    for target, amount in _investment_balances(values, cfg).items():
        found[target][through.strftime("%Y-%m")] = (amount, through.isoformat())
    # Where bank exports are the record, the sheet's running balances (typed, unverified) don't apply.
    for name, spans in _bank_windows(conn).items():
        first = min(start for start, _ in spans)
        for target in [t for t in found if t[0] == name]:
            found[target] = {m: v for m, v in found[target].items() if v[1] < first}
    result["balances"] = _write_balances(conn, found)
    result["loans_added"] = _write_loans(conn, loans)
    conn.commit()
    return result
